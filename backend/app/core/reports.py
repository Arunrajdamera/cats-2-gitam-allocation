from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from .schemas import AllocationRun


def build_statistics(run: AllocationRun) -> dict[str, float | int]:
    allocations = [allocation for result in run.faculty for allocation in result.allocations]
    total_sections = sum(subject.capacity for subject in run.subjects.values())
    allocated_sections = total_sections - sum(run.remaining_capacity.values())
    priority_counts = {f"priority_{i}": sum(1 for a in allocations if a.priority == i) for i in range(1, 5)}
    p1_faculty = sum(1 for result in run.faculty if any(allocation.priority == 1 for allocation in result.allocations))
    two_slot = 0
    same_subject = 0
    mixed_subject = 0
    for result in run.faculty:
        normal = sorted([item for item in result.allocations if item.slot <= 2], key=lambda item: item.slot)
        if len(normal) == 2:
            two_slot += 1
            if normal[0].subject_code == normal[1].subject_code:
                same_subject += 1
            else:
                mixed_subject += 1
    workloads = [result.workload_hours for result in run.faculty]
    avg_workload = sum(workloads) / len(workloads) if workloads else 0
    variance = sum((value - avg_workload) ** 2 for value in workloads) / len(workloads) if workloads else 0
    return {
        "total_faculty": len(run.faculty),
        "total_subjects": len(run.subjects),
        "total_sections": total_sections,
        "allocated_sections": allocated_sections,
        "remaining_sections": sum(run.remaining_capacity.values()),
        "subject_fulfillment_percent": round((allocated_sections / total_sections) * 100, 2) if total_sections else 0,
        "coverage_percent": round((allocated_sections / total_sections) * 100, 2) if total_sections else 0,
        "coverage_allocations": sum(1 for a in allocations if a.status == "COVERAGE"),
        "related_match_allocations": sum(1 for a in allocations if a.priority == "related"),
        **priority_counts,
        "priority_1_allocation_percent": round((priority_counts["priority_1"] / max(len(allocations), 1)) * 100, 2),
        "priority_1_faculty_satisfaction_percent": round((p1_faculty / max(len(run.faculty), 1)) * 100, 2),
        "same_subject_percent": round((same_subject / max(two_slot, 1)) * 100, 2),
        "mixed_subject_percent": round((mixed_subject / max(two_slot, 1)) * 100, 2),
        "manual_allocation_count": sum(1 for result in run.faculty if "MANUAL" in result.status),
        "unallocated_count": sum(1 for result in run.faculty if not result.allocations),
        "average_workload": round(avg_workload, 2),
        "max_workload": max(workloads) if workloads else 0,
        "min_workload": min(workloads) if workloads else 0,
        "workload_variance": round(variance, 2),
        "validation_passed": sum(1 for i in run.validation if i.status == "PASS"),
        "validation_warnings": sum(1 for i in run.validation if i.status == "WARNING"),
        "validation_failures": sum(1 for i in run.validation if i.status == "FAIL"),
        "runtime_seconds": round(run.runtime_seconds, 3),
    }


def export_excel(run: AllocationRun, path: str | Path) -> Path:
    path = Path(path)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Final Faculty Allocation"
    headers = [
        "S.No",
        "EMP ID",
        "Faculty Name",
        "Designation",
        "Admin Role",
        "Priority 1 Preference",
        "Priority 2 Preference",
        "Priority 3 Preference",
        "Priority 4 Preference",
        "Priority 5 Preference",
        "Priority 6 Preference",
        "Priority 7 Preference",
        "Priority 8 Preference",
        "Subject 1 Allocated",
        "P1 Used",
        "Subject 2 Allocated",
        "P2 Used",
        "Total Sub",
        "WL hrs",
        "Same Sub?",
        "Allocation Status",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
    for index, result in enumerate(run.faculty, start=1):
        normal_allocations = sorted([item for item in result.allocations if item.slot <= 2], key=lambda item: item.slot)
        first = normal_allocations[0] if normal_allocations else None
        second = normal_allocations[1] if len(normal_allocations) > 1 else None
        prefs = [
            f"{code} - {run.subjects[code].name}" if code and code in run.subjects else ""
            for code in result.preferences[:8]
        ]
        same_subject = "YES" if first and second and first.subject_code == second.subject_code else ""
        ws.append(
            [
                index,
                result.emp_id,
                result.faculty_name,
                result.designation,
                result.admin_role,
                *prefs,
                f"{first.subject_code} - {first.subject_name}" if first else "",
                first.priority if first else "",
                f"{second.subject_code} - {second.subject_name}" if second else "",
                second.priority if second else "",
                len(normal_allocations),
                result.workload_hours,
                same_subject,
                result.status,
            ]
        )

    ws2 = wb.create_sheet("Subject Utilization")
    ws2.append([
        "Subject Code",
        "Subject Name",
        "Semester",
        "Category",
        "L",
        "P",
        "Credits",
        "Weekly Classes",
        "Required Sections",
        "Allocated Sections",
        "Remaining Sections",
        "Faculty Requirement",
        "Instructor Workload",
        "Partner Workload",
        "Fulfillment Status",
    ])
    for code, subject in run.subjects.items():
        remaining = run.remaining_capacity.get(code, 0)
        ws2.append([
            code,
            subject.name,
            subject.semester,
            subject.category,
            subject.l,
            subject.p,
            subject.credits,
            subject.weekly_classes,
            subject.capacity,
            subject.capacity - remaining,
            remaining,
            subject.faculty_requirement,
            subject.instructor_workload,
            subject.partner_workload,
            "FULL" if remaining == 0 else "PARTIAL",
        ])

    ws3 = wb.create_sheet("Manual Allocation Required")
    ws3.append(["EMP ID", "Faculty Name", "Status"])
    for result in run.faculty:
        if "MANUAL" in result.status or "PARTIALLY" in result.status:
            ws3.append([result.emp_id, result.faculty_name, result.status])

    ws4 = wb.create_sheet("Coverage Allocations")
    ws4.append(["EMP ID", "Faculty Name", "Designation", "Admin Role", "Subject Code", "Subject Name", "Slot", "Priority", "Status", "Workload"])
    for result in run.faculty:
        for allocation in result.allocations:
            if allocation.status == "COVERAGE":
                ws4.append([result.emp_id, result.faculty_name, result.designation, result.admin_role, allocation.subject_code, allocation.subject_name, allocation.slot, allocation.priority, allocation.status, result.workload_hours])

    ws5 = wb.create_sheet("Validation Report")
    ws5.append(["Check", "Status", "Message", "Row", "Code"])
    for issue in run.validation:
        ws5.append([issue.check, issue.status, issue.message, issue.row, issue.code])

    ws6 = wb.create_sheet("Statistics")
    for key, value in build_statistics(run).items():
        ws6.append([key, value])

    ws7 = wb.create_sheet("Audit Log")
    ws7.append(["Timestamp", "Admin User", "Action Type", "Reason", "Before", "After"])
    for event in run.audit_log:
        ws7.append([event.timestamp.isoformat(), event.admin_user, event.action_type, event.reason, str(event.before_value), str(event.after_value)])

    ws8 = wb.create_sheet("Manual Corrections")
    ws8.append(["EMP ID", "New Subject", "Reason", "Status"])

    ws9 = wb.create_sheet("Workload Report")
    ws9.append(["EMP ID", "Faculty Name", "Designation", "Workload Hours", "Allocation Count", "Status"])
    for result in run.faculty:
        ws9.append([result.emp_id, result.faculty_name, result.designation, result.workload_hours, len(result.allocations), result.status])

    ws10 = wb.create_sheet("Priority Satisfaction")
    ws10.append(["Priority", "Allocation Count"])
    for key, value in build_statistics(run).items():
        if key.startswith("priority_") and key.endswith(tuple(str(i) for i in range(1, 5))):
            ws10.append([key, value])
    wb.save(path)
    return path


def export_pdf_summary(run: AllocationRun, path: str | Path) -> Path:
    path = Path(path)
    stats = build_statistics(run)
    doc = canvas.Canvas(str(path), pagesize=A4)
    y = 800
    doc.setFont("Helvetica-Bold", 14)
    doc.drawString(40, y, "GITAM CSE Faculty Subject Allocation Summary")
    y -= 30
    doc.setFont("Helvetica", 10)
    for key, value in stats.items():
        doc.drawString(40, y, f"{key.replace('_', ' ').title()}: {value}")
        y -= 16
        if y < 60:
            doc.showPage()
            y = 800
            doc.setFont("Helvetica", 10)
    doc.save()
    return path
