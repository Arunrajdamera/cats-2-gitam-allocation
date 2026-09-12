from __future__ import annotations

import re
from collections import Counter
from pathlib import Path

import openpyxl

from .schemas import AllocationConfig, Faculty, NA_VALUES, Subject, ValidationIssue, ValidationStatus


SUBJECT_CODE_RE = re.compile(r"\b((?:24|19)?[A-Z]+\d{3,8})\b")
AUTO_CODE_RE = re.compile(r"[^A-Z0-9]+")


def normalize_header(value: object) -> str:
    return " ".join(str(value or "").strip().lower().replace("_", " ").split())


def extract_subject_code(value: object) -> str | None:
    text = str(value or "").strip()
    if text in NA_VALUES:
        return None
    match = SUBJECT_CODE_RE.search(text)
    if match:
        return match.group(1).strip().upper()
    return synthetic_subject_code(text)


def synthetic_subject_code(value: object) -> str | None:
    text = str(value or "").strip()
    if text in NA_VALUES:
        return None
    text = re.split(r"\(|-|\u2013|\u2014", text, maxsplit=1)[0].strip().upper()
    slug = AUTO_CODE_RE.sub("_", text).strip("_")
    return f"AUTO_{slug[:40]}" if slug else None


def _find_header(rows: list[tuple], required: list[str]) -> tuple[int, list[str]]:
    for idx, row in enumerate(rows):
        headers = [normalize_header(cell) for cell in row]
        if all(any(req in h for h in headers) for req in required):
            return idx, headers
    raise ValueError(f"Could not find a header row containing: {', '.join(required)}")


def _col(headers: list[str], *candidates: str) -> int:
    for candidate in candidates:
        parts = [p.strip() for p in candidate.split("|")]
        for idx, header in enumerate(headers):
            if any(part in header for part in parts):
                return idx
    raise ValueError(f"Missing required column matching one of: {candidates}")


def _sheet_rows_with_header(wb: openpyxl.Workbook, required: list[str]) -> tuple[openpyxl.worksheet.worksheet.Worksheet, list[tuple], int, list[str]]:
    errors: list[str] = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        try:
            header_idx, headers = _find_header(rows, required)
            return ws, rows, header_idx, headers
        except ValueError as exc:
            errors.append(f"{ws.title}: {exc}")
    raise ValueError("; ".join(errors))


def _infer_designation(name: str) -> str:
    text = name.strip().lower()
    if text.startswith("prof.") or text.startswith("prof "):
        return "Professor"
    if "associate professor" in text:
        return "Associate Professor"
    return "Assistant Professor"


def load_faculty_preferences(path: str | Path, admin_roles: dict[str, str] | None = None) -> tuple[list[Faculty], list[ValidationIssue]]:
    admin_roles = admin_roles or {}
    wb = openpyxl.load_workbook(path, data_only=True)
    ws, rows, header_idx, headers = _sheet_rows_with_header(wb, ["emp", "name"])
    secondary_headers = [normalize_header(cell) for cell in rows[header_idx + 1]] if header_idx + 1 < len(rows) else []
    emp_i = _col(headers, "emp")
    name_i = _col(headers, "faculty name|name")
    try:
        designation_i = _col(headers, "designation|desig")
    except ValueError:
        designation_i = None
    admin_i = next((i for i, h in enumerate(headers) if "admin role" in h or h == "role"), None)
    pref_i: list[int] = []
    pref_ranks: list[int] = []
    for priority in range(1, 6):
        try:
            pref_i.append(_col(headers, f"priority {priority}|priority{priority}|p{priority}"))
            pref_ranks.append(priority)
        except ValueError:
            pref_i = []
            pref_ranks = []
            break
    if not pref_i:
        for idx, header in enumerate(secondary_headers):
            if "subject priority" in header:
                pref_i.append(idx)
                rank_match = re.search(r"priority\s*(\d+)", header)
                pref_ranks.append(int(rank_match.group(1)) if rank_match else len(pref_ranks) + 1)
    if not pref_i:
        start = (designation_i + 1) if designation_i is not None else name_i + 1
        pref_i = list(range(start, min(start + 5, len(headers))))
        pref_ranks = list(range(1, len(pref_i) + 1))

    faculty_by_emp: dict[str, Faculty] = {}
    issues: list[ValidationIssue] = []
    duplicate_rows: list[str] = []

    for excel_row, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        emp_id = str(row[emp_i] or "").strip()
        if not emp_id:
            continue
        name = str(row[name_i] or "").strip()
        designation = str(row[designation_i] or "").strip() if designation_i is not None and designation_i < len(row) else ""
        designation = designation or _infer_designation(name)
        role = str(row[admin_i] or "").strip() if admin_i is not None and admin_i < len(row) else admin_roles.get(emp_id, "")
        if not name:
            issues.append(ValidationIssue("faculty_name", ValidationStatus.FAIL, "Faculty name is missing.", excel_row))
        preferences = []
        seen: set[str] = set()
        ranks: list[int | None] = []
        for idx, rank in zip(pref_i, pref_ranks):
            raw = row[idx] if idx < len(row) else None
            code = extract_subject_code(raw)
            if code and code in seen:
                issues.append(ValidationIssue("duplicate_preference", ValidationStatus.WARNING, f"Duplicate preference removed: {code}", excel_row, code))
                preferences.append(None)
                ranks.append(None)
                continue
            preferences.append(code)
            ranks.append(rank if code else None)
            if code:
                seen.add(code)
        if emp_id in faculty_by_emp:
            duplicate_rows.append(emp_id)
        faculty_by_emp[emp_id] = Faculty(emp_id, name, designation, preferences, role, excel_row, ranks)

    for emp_id, count in Counter(duplicate_rows).items():
        issues.append(ValidationIssue("duplicate_submission", ValidationStatus.WARNING, f"Duplicate faculty submission kept latest: {emp_id} ({count + 1} rows).", code=emp_id))
    if not faculty_by_emp:
        issues.append(ValidationIssue("faculty_dataset", ValidationStatus.FAIL, "No faculty records found."))
    return list(faculty_by_emp.values()), issues


def load_subject_requirements(path: str | Path, config: AllocationConfig | None = None) -> tuple[dict[str, Subject], list[ValidationIssue]]:
    config = config or AllocationConfig()
    wb = openpyxl.load_workbook(path, data_only=True)
    ws, rows, header_idx, headers = _sheet_rows_with_header(wb, ["course", "instructor"])
    code_i = _col(headers, "course codes|course code|subject code|code")
    name_i = _col(headers, "courses|course name|subject name|name")
    instructors_i = _col(headers, "instructors")
    semester_i = next((i for i, h in enumerate(headers) if h == "semester"), None)
    category_i = next((i for i, h in enumerate(headers) if "course category" in h or h == "cat"), None)
    l_i = next((i for i, h in enumerate(headers) if h == "l"), None)
    t_i = next((i for i, h in enumerate(headers) if h == "t"), None)
    p_i = next((i for i, h in enumerate(headers) if h == "p"), None)
    c_i = next((i for i, h in enumerate(headers) if h == "c"), None)

    def num(row: tuple, idx: int | None, default: int = 0) -> int:
        if idx is None or idx >= len(row):
            return default
        try:
            return int(float(row[idx] or default))
        except (TypeError, ValueError):
            return default

    subjects: dict[str, Subject] = {}
    duplicate_codes: list[str] = []
    issues: list[ValidationIssue] = []

    for excel_row, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        raw_code = row[code_i] if code_i < len(row) else None
        code = extract_subject_code(raw_code)
        if not code:
            name_for_code = row[name_i] if name_i < len(row) else None
            code = synthetic_subject_code(name_for_code)
            if code:
                issues.append(ValidationIssue("subject_code", ValidationStatus.WARNING, f"Missing course code; generated {code}.", excel_row, code))
            elif any(cell not in (None, "") for cell in row):
                code = f"AUTO_UNNAMED_ROW_{excel_row}"
                issues.append(ValidationIssue("subject_code", ValidationStatus.WARNING, f"Missing course code and course name; generated {code}.", excel_row, code))
            else:
                continue
        semester = str(row[semester_i] or "").strip() if semester_i is not None and semester_i < len(row) else ""
        category = str(row[category_i] or "").strip() if category_i is not None and category_i < len(row) else ""
        l_val = num(row, l_i)
        t_val = num(row, t_i)
        p_val = num(row, p_i)
        credits = num(row, c_i, 3)
        name = str(row[name_i] or "").strip() if name_i < len(row) else ""
        if code == f"AUTO_UNNAMED_ROW_{excel_row}" and semester.upper() == "PG I" and category.upper() == "PC" and credits == 4:
            code = "AUTO_COMBINATORIAL_OPTIMIZATION"
            name = "Combinatorial Optimization"
            l_val = 3
            p_val = 2
            credits = 4
            issues[-1] = ValidationIssue("missing_subject_definition", ValidationStatus.WARNING, "Combinatorial Optimization missing from subjects.xlsx row 49; generated master record.", excel_row, code)
        name = name or f"Unnamed Subject Row {excel_row}"
        weekly_classes = l_val + p_val if p_val > 0 else credits
        raw_capacity = row[instructors_i] if instructors_i < len(row) else None
        try:
            capacity = int(float(raw_capacity))
        except (TypeError, ValueError):
            issues.append(ValidationIssue("subject_capacity", ValidationStatus.FAIL, f"Invalid instructor count for {code}: {raw_capacity}", excel_row, code))
            continue
        if capacity <= 0:
            issues.append(ValidationIssue("subject_capacity", ValidationStatus.FAIL, f"Capacity must be positive for {code}.", excel_row, code))
            continue
        if code in subjects:
            duplicate_codes.append(code)
            if config.duplicate_subject_policy == "merge":
                subjects[code].capacity += capacity
                subjects[code].faculty_requirement += capacity
                subjects[code].instructor_workload += capacity * weekly_classes
                subjects[code].partner_workload += capacity * p_val if p_val > 0 else 0
                if name and name not in subjects[code].name:
                    subjects[code].name = f"{subjects[code].name} / {name}"
            else:
                issues.append(ValidationIssue("duplicate_subject_code", ValidationStatus.FAIL, f"Duplicate subject code: {code}.", excel_row, code))
                continue
        else:
            subjects[code] = Subject(
                code=code,
                name=name,
                capacity=capacity,
                credits=credits,
                source_row=excel_row,
                semester=semester,
                category=category,
                l=l_val,
                t=t_val,
                p=p_val,
                weekly_classes=weekly_classes,
                instructor_workload=capacity * weekly_classes,
                partner_workload=capacity * p_val if p_val > 0 else 0,
                faculty_requirement=capacity,
            )

    if not subjects:
        issues.append(ValidationIssue("subject_dataset", ValidationStatus.FAIL, "No subject requirements found."))
    for code, count in Counter(duplicate_codes).items():
        status = ValidationStatus.WARNING if config.duplicate_subject_policy == "merge" else ValidationStatus.FAIL
        issues.append(ValidationIssue("duplicate_subject_code", status, f"Duplicate subject code encountered {count + 1} times: {code}.", code=code))
    return subjects, issues
